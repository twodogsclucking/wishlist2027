# -*- coding: utf-8 -*-
"""ONE-TIME migration: Wishlist.xlsx -> data/items.json + data/settings.json.

After this runs once, data/items.json is the master. Edit the JSON, not the workbook.
Kept in the repo only as a record of how the data was first produced.
NOTE: it writes the old flat "freight" field. items.json has since moved to origin/shipping/basis
(landed-cost model), so do not re-run it over the current data.
"""
import json, os, sys
import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
wb = openpyxl.load_workbook(os.path.join(ROOT, "Wishlist.xlsx"))
wsS, wsO, wsA = wb["Wishlist Summary"], wb["Vehicle & Equipment Options"], wb["Activity Assumptions"]

def T(mn, en): return {"mn": mn, "en": en}

# ------------------------------------------------------------------ read the workbook
LINES = {}                                     # summary item name -> row data
for r in wsS.iter_rows(min_row=6, max_row=41):
    if r[0].value and r[2].value is not None:
        LINES[r[0].value] = dict(qty=r[1].value, unit=r[2].value, freight=r[3].value or 0,
                                 phase=r[5].value, paid=r[6].value)

GROUPS, cur = {}, None                         # option group title -> {options, selected, qty}
for row in wsO.iter_rows(min_row=4, max_row=72):
    a, b = row[0].value, row[1].value
    if a and b is None and row[2].value is None and not str(a).startswith(("->", "Quantity")):
        cur = a; GROUPS[cur] = {"options": {}, "selected": None, "qty": None}
    elif a == "Option": pass
    elif a and str(a).startswith("->"): GROUPS[cur]["selected"] = b
    elif a and str(a).startswith("Quantity"): GROUPS[cur]["qty"] = b
    elif a and isinstance(b, (int, float)) and cur: GROUPS[cur]["options"][a] = b

hdrs = [r for r in range(1, 90) if wsA.cell(r, 1).value == "Activity"]
def block(h):
    out, r = {}, h + 1
    while wsA.cell(r, 1).value and wsA.cell(r, 2).value is not None:
        out[wsA.cell(r, 1).value] = [wsA.cell(r, c).value for c in (2, 3, 4)]   # cons, base, aggr
        r += 1
    return out
PRICE, FREQ, SEAS, OPEX = (block(h) for h in hdrs[:4])
GUEST_NIGHTS = wsA.cell(5, 2).value

def scenarios(xl_name):
    keys = ("cons", "base", "aggr")
    return {k: {"price": PRICE[xl_name][i], "perDay": FREQ[xl_name][i],
                "days": SEAS[xl_name][i], "opex": OPEX[xl_name][i]} for i, k in enumerate(keys)}

# ------------------------------------------------------------------ helpers
def usd(a): return {"amount": round(a, 2), "currency": "USD"}
def mnt(a): return {"amount": int(round(a)), "currency": "MNT"}

def fixed(xl, name, src, unit=None, freight=None, price=None):
    """A component with a single option, read from one summary line."""
    L = LINES[xl]
    c = {"id": "main", "quantity": L["qty"], "options": [{
        "id": "main", "name": name, "source": src,
        "price": price or usd(L["unit"]),
        "freight": L["freight"] if freight is None else freight, "default": True}]}
    if unit: c["unit"] = unit
    return c

def swap(group, qty_unit=None, opts=()):
    """A swappable component. opts = (excel option name, id, name, note, source, freight[, price])."""
    G = GROUPS[group]
    out = []
    for o in opts:
        xl, oid, nm, note, src, fr = o[:6]
        price = o[6] if len(o) > 6 else usd(G["options"][xl])
        out.append({"id": oid, "name": nm, "note": note, "source": src, "price": price,
                    "freight": fr, "default": xl == G["selected"]})
    c = {"id": "main", "quantity": G["qty"], "options": out}
    if qty_unit: c["unit"] = qty_unit
    return c

M2, M, PCS = T("м²", "m²"), T("м", "m"), None

# ------------------------------------------------------------------ the 28 items
I = []

I.append(dict(id="storage-parking", category="buildings", phase=1, earns="partial",
  name=T("Дулаан зогсоол, агуулах", "Heated parking & storage"),
  summary=T("5 машины дулаан зогсоолтой, жагсаалтын бүх тээврийн хэрэгсэл, тоног төхөөрөмжийг өвөлжүүлэх дулаалгатай барилга. Зочдын зогсоолыг төлбөртэй үйлчилгээ болгож болно.",
            "A heated, insulated building with parking for 5 cars and winter storage for every vehicle on this list. Guest parking can be sold as a paid service."),
  components=[fixed("Storage & Parking Building (5 cars + equipment)", T("Барилга", "Building"), "construction-benchmark", unit=M2)],
  revenue=dict(unit=T("машин/шөнө", "per car-night"), attachApplies=False, capexShare=0.5,
               scenarios=scenarios("Storage & Parking Building (5 cars + equipment)"))))

hangar = [fixed("Helicopter Hangar (structure incl. large door)", T("Ангарын барилга", "Hangar building"), "construction-benchmark", unit=M2),
          fixed("Helicopter Hangar - internal fuel tank", T("Түлшний сав", "Fuel tank"), "research-global"),
          fixed("Helicopter Hangar - electric ground dolly/tug", T("Нисдэг тэрэг зөөгч", "Electric tug"), "research-global")]
for c, cid, nm in zip(hangar, ("building", "fuel-tank", "tug"),
                      (T("Ангарын барилга", "Hangar building"), T("Түлшний сав", "Fuel tank"), T("Нисдэг тэрэг зөөгч", "Electric tug"))):
    c["id"] = cid; c["name"] = nm
I.append(dict(id="helicopter-hangar", category="buildings", phase=3, earns="no",
  name=T("Нисдэг тэрэгний ангар", "Helicopter hangar"),
  summary=T("Өвөл нисдэг тэрэг хонох боломжтой дулаан ангар. Түлшний сав, нисдэг тэрэг зөөх цахилгаан тэрэгтэй. Шууд орлого олохгүй.",
            "A heated hangar so visiting helicopters can stay overnight in winter, with an on-site fuel tank and an electric tug. It earns no direct revenue."),
  components=hangar))

I.append(dict(id="river-dock", category="buildings", phase=1, earns="no",
  name=T("Усан онгоцны зогсоол", "River dock"),
  summary=T("Зочид усан онгоцонд аюулгүй суух, буух голын эргийн тавцан.", "A boarding platform so guests can get on and off the jet boat safely."),
  components=[fixed("River Dock (guest boarding platform)", T("Тавцан", "Platform"), "construction-benchmark", unit=M2)]))

I.append(dict(id="river-crossing", category="buildings", phase=1, earns="no",
  name=T("Гол гарам", "River crossing"),
  summary=T("Жижиг горхи дээгүүр гарах гарам (ердийн үед ~1м, үерийн үед ~10м). Хоолойт гарам машин явах боломжтой бөгөөд гүүрээс хамаагүй хямд.",
            "A crossing over the small stream (about 1 m wide, up to 10 m in flood). A culvert takes cars and costs far less than a bridge."),
  components=[swap("RIVER CROSSING - stream approx 1m wide, up to approx 10m at flood", opts=(
    ("Pipe culvert battery + embankment crossing", "pipe-culvert", T("Хоолойт гарам, далантай", "Pipe culvert with embankment"),
     T("Машин явах хамгийн хямд сонголт — шороон дор 2–3 том хоолой.", "Cheapest car-capable option — 2–3 large pipes under fill."), "construction-benchmark", 0),
    ("Precast box culvert crossing", "box-culvert", T("Бэлэн бетон хайрцган гарам", "Precast box culvert"),
     T("10м үерийн өргөнийг илүү сайн даах, голын ёроолыг хэвээр нь үлдээнэ.", "Copes better with the 10 m flood width and keeps the streambed natural."), "construction-benchmark", 0),
    ("Modular steel clear-span vehicle bridge", "steel-bridge", T("Угсармал ган гүүр", "Modular steel bridge"),
     T("Бүтэн алгасалттай гүүр — 1м горхинд хэтэрхий их байж магадгүй.", "A true clear span — likely more than a 1 m stream needs."), "construction-benchmark", 0)))]))

I.append(dict(id="boardwalk", category="buildings", phase=2, earns="no",
  name=T("Модон зам", "Wooden boardwalk"),
  summary=T("Зочдыг гол руу хялбар хүргэх 100 метр модон зам.", "A 100 m wooden walkway giving guests easy access to the river."),
  components=[fixed("Wooden Boardwalk (approx 100m, river access)", T("Зам", "Walkway"), "construction-benchmark", unit=M)]))

I.append(dict(id="concrete-path", category="buildings", phase=2, earns="no",
  name=T("Бетон зам", "Concrete path"),
  summary=T("Клуб байрнаас хог хаягдлын цэг хүртэлх 50 метр бетон зам — жилийн турш ашиглана.", "A 50 m concrete path from the clubhouse to the waste point, usable all year."),
  components=[fixed("Concrete Path (clubhouse to waste point, 50m)", T("Зам", "Path"), "construction-benchmark", unit=M2)]))

I.append(dict(id="basketball-court", category="buildings", phase=2, earns="no",
  name=T("3x3 сагсан бөмбөгийн талбай", "3x3 basketball court"),
  summary=T("Зочид болон ажилчдад зориулсан гадна 3x3 талбай.", "An outdoor 3x3 court for guests and staff."),
  components=[fixed("3x3 Outdoor Basketball Court", T("Талбай", "Court"), "research-global")]))

I.append(dict(id="viewpoint-rest-area", category="buildings", phase=2, earns="no",
  name=T("Уулын ажиглалтын цэг", "Mountain viewpoint rest area"),
  summary=T("Уулын үзэсгэлэнт цэг дээрх модон ширээ, амрах талбай.", "A wooden table and rest area at a mountain viewpoint."),
  components=[fixed("Wooden Table + Rest Area (mountain viewpoint)", T("Амрах талбай", "Rest area"), "construction-benchmark")]))

I.append(dict(id="telecom-station", category="buildings", phase=2, earns="no",
  name=T("Холбооны станц", "Telecom base station"),
  summary=T("Лоджид тогтвортой утас, интернэтийн сүлжээ өгөх холбооны станц.", "A mobile base station giving the lodge reliable phone and data coverage."),
  components=[fixed("Telecom Base Station", T("Станц", "Station"), "given", price=mnt(350000000), freight=0)]))

# ---- vehicles
I.append(dict(id="golf-carts", category="vehicles", phase=1, earns="yes",
  name=T("Цахилгаан гольфын тэрэг", "Electric golf carts"),
  summary=T("Зочдын аялал, ачаа тээшинд ашиглах цахилгаан тэрэг.", "Electric carts for guest tours and moving luggage around the property."),
  components=[swap("GOLF CART - electric", opts=(
    ("Basic 4-seat electric, lead-acid battery", "basic-lead-acid", T("Энгийн 4 суудалтай, хар тугалган батарей", "Basic 4-seat, lead-acid battery"),
     T("Хамгийн хямд ч батарейны насжилт богино.", "Cheapest, but shorter battery life."), "research-china", 0.20),
    ("Premium 4-seat electric, lithium battery", "premium-lithium", T("Сайн чанарын 4 суудалтай, лити батарей", "Premium 4-seat, lithium battery"),
     T("Өдөр бүрийн аялалд илүү хол явж, батарей удаан эдэлнэ.", "Longer range and battery life for daily tours."), "research-china", 0.20)))],
  revenue=dict(unit=T("хагас өдрийн түрээс", "per half-day rental"), attachApplies=True, capexShare=1.0, scenarios=scenarios("Electric Golf Carts"))))

I.append(dict(id="utv", category="vehicles", phase=2, earns="yes",
  name=T("UTV машин", "UTVs"),
  summary=T("Хөтөчтэй аялалд зориулсан бартаат замын UTV машин.", "Side-by-side off-road vehicles for guided tours."),
  components=[swap("UTV / SIDE-BY-SIDE", opts=(
    ("Standard 4x4 UTV - gas - 4-seat (generic China brand)", "generic-petrol", T("Энгийн 4x4, бензин, 4 суудал", "Standard 4x4, petrol, 4 seats"),
     T("Хятадын энгийн брэндийн хямд сонголт.", "Budget option from a generic Chinese brand."), "research-china", 0.20),
    ("CFMOTO UFORCE-series (established brand)", "cfmoto-uforce", T("CFMOTO UFORCE", "CFMOTO UFORCE"),
     T("Нэр хүндтэй брэнд — найдвартай, сэлбэг олдоцтой.", "Established brand — better reliability and spare parts."), "research-china", 0.20),
    ("Electric UTV 4-seat", "electric-utv", T("Цахилгаан UTV, 4 суудал", "Electric UTV, 4 seats"),
     T("Лоджийн батарей, цэнэглэгчтэй зохицно.", "Pairs with the lodge's battery and charger."), "research-china", 0.20)))],
  revenue=dict(unit=T("хагас өдрийн түрээс", "per half-day rental"), attachApplies=True, capexShare=1.0, scenarios=scenarios("UTV"))))

I.append(dict(id="snowmobiles", category="vehicles", phase=2, earns="yes",
  name=T("Цасны мотоцикл", "Snowmobiles"),
  summary=T("Өвлийн аялалд зориулсан цасны мотоцикл. Цахилгааныг сонгосон — -20°C-аас доош хүйтэнд батарей хурдан суларна.",
            "Snowmobiles for winter guest tours. Electric was chosen — note that battery range drops sharply below -20°C."),
  components=[swap("SNOWMOBILE - electric requested", opts=(
    ("Taiga Nomad - fully electric", "taiga-nomad", T("Taiga Nomad — бүрэн цахилгаан", "Taiga Nomad — fully electric"),
     T("Хүссэнээр цахилгаан. Бензинийхээс ~4 дахин үнэтэй; хүйтэнд явах зай эрс багасна.", "Electric, as requested. About 4× the price of petrol; range drops sharply in deep cold."), "research-global", 0.30),
    ("Entry-level gas snowmobile approx 200cc", "petrol-200cc", T("Энгийн бензин, ~200cc", "Entry-level petrol, ~200cc"),
     T("Харьцуулахад зориулсан — хамаагүй хямд, хүйтний эрсдэлгүй.", "Shown for comparison — far cheaper, no cold-weather battery risk."), "research-china", 0.30)))],
  revenue=dict(unit=T("хагас өдрийн түрээс", "per half-day rental"), attachApplies=True, capexShare=1.0, scenarios=scenarios("Snowmobile (electric)"))))

I.append(dict(id="snow-scooters", category="vehicles", phase=2, earns="yes",
  name=T("Цасны скүтер", "Snow scooters"),
  summary=T("Өвлийн богино түрээсэд зориулсан, зогсож явдаг нэг хүний цасны скүтер.", "Stand-up, single-rider snow scooters for short winter rentals."),
  components=[swap("SNOW SCOOTER - standing, single person", opts=(
    ("Budget stand-up electric snow scooter", "budget", T("Хямд скүтер", "Budget scooter"),
     T("Хүч, бат бөх чанар бага.", "Lower power and durability."), "research-china", 0.20),
    ("Performance stand-up electric snow scooter", "performance", T("Хүчирхэг скүтер", "Performance scooter"),
     T("Байнгын түрээсэд тэсвэртэй.", "Built for repeated rental use."), "research-china", 0.20)))],
  revenue=dict(unit=T("нэг түрээс", "per rental"), attachApplies=True, capexShare=1.0, scenarios=scenarios("Snow Scooter (standing, single person)"))))

I.append(dict(id="surron-bikes", category="vehicles", phase=1, earns="yes",
  name=T("Surron цахилгаан мотоцикл", "Surron e-bikes"),
  summary=T("Зуны замын түрээсэд зориулсан цахилгаан мотоцикл.", "Electric dirt bikes for summer trail rentals."),
  components=[swap("SURRON ELECTRIC BIKE", opts=(
    ("Surron Light Bee X - standard import", "standard", T("Surron Light Bee X — энгийн импорт", "Surron Light Bee X — standard import"),
     T("Үйлдвэрийн ердийн үнэ.", "Standard factory pricing."), "research-china", 0.20),
    ("Surron Light Bee X - premium market retail", "premium", T("Surron Light Bee X — дээд үнийн сегмент", "Surron Light Bee X — premium retail"),
     T("Жижиглэнгийн өндөр үнэ.", "Higher retail price point."), "research-global", 0.20)))],
  revenue=dict(unit=T("хагас өдрийн түрээс", "per half-day rental"), attachApplies=True, capexShare=1.0, scenarios=scenarios("Surron Electric Bikes"))))

I.append(dict(id="jet-boat", category="vehicles", phase=1, earns="yes",
  name=T("Усан онгоц (жет)", "Jet boat"),
  summary=T("Голын аялалд зориулсан жет усан онгоц. Үнэ нь дилерийн бодит үнийн санал.", "A jet boat for guided river rides. Prices are real dealer quotations."),
  components=[swap("JET BOAT", opts=(
    ("Fregat 480 FM Jet - 2-stroke 40hp - 2 to 4 seats", "fregat-480-jet", T("Fregat 480 FM Jet — 40 морь, 2–4 суудал", "Fregat 480 FM Jet — 40 hp, 2–4 seats"),
     T("Хамгийн жижиг, хямд; 2 тактын хөдөлгүүр.", "Smallest and cheapest; 2-stroke engine."), "quote", 0, mnt(65900000)),
    ("Fregat 480 FM Jet SP - 4-stroke 70hp - 2 to 4 seats", "fregat-480-jet-sp", T("Fregat 480 FM Jet SP — 70 морь, 2–4 суудал", "Fregat 480 FM Jet SP — 70 hp, 2–4 seats"),
     T("Ижил хэмжээтэй, 4 тактын илүү найдвартай хөдөлгүүр.", "Same size, quieter and more reliable 4-stroke engine."), "quote", 0, mnt(145000000)),
    ("Fregat F600 Jet SP - 4-stroke 115hp - 6 to 10 seats", "fregat-f600-jet-sp", T("Fregat F600 Jet SP — 115 морь, 6–10 суудал", "Fregat F600 Jet SP — 115 hp, 6–10 seats"),
     T("Нэг удаад 6–10 зочин — бүлгийн аялалд хамгийн тохиромжтой.", "Carries 6–10 guests a trip — best for group tours."), "quote", 0, mnt(210000000))))],
  revenue=dict(unit=T("зочин бүрээр, нэг аялал", "per guest, per ride"), attachApplies=True, capexShare=1.0, scenarios=scenarios("Jet Boat"))))

I.append(dict(id="horse-carriage", category="vehicles", phase=2, earns="yes",
  name=T("Морин тэрэг", "Horse-drawn carriage"),
  summary=T("Монгол шороон замд тохирсон, 4-өөс дээш хүний суудалтай бат бөх морин тэрэг.", "A sturdy carriage seating 4+, built for Mongolian dirt roads, for paid rides."),
  components=[swap("HORSE-DRAWN CARRIAGE", opts=(
    ("Locally built Mongolian carriage - sturdy, 4+ seats", "local-built", T("Монголд хийсэн", "Built locally in Mongolia"),
     T("Хамаагүй хямд, орон нутгийн шороон замд зориулсан.", "Far cheaper and made for local dirt roads."), "research-mongolia", 0),
    ("Imported custom-built carriage", "imported", T("Гадаадаас захиалсан", "Imported custom carriage"),
     T("Өндөр үнэ, том ачааны тээврийн зардалтай.", "Higher cost plus oversized freight."), "research-global", 0.30)))],
  revenue=dict(unit=T("нэг аялал", "per ride"), attachApplies=True, capexShare=1.0, scenarios=scenarios("Horse-Drawn Carriage"))))

# ---- safety & wildlife
I.append(dict(id="thermal-drone", category="safety", phase=2, earns="yes",
  name=T("Дулааны дрон", "Thermal drone"),
  summary=T("Шөнийн ан амьтан ажиглах хөтөчтэй аялалд зориулсан дулааны камертай дрон.", "A thermal-camera drone for guided night-time wildlife spotting."),
  components=[fixed("Thermal Drone (wildlife spotting)", T("Дрон", "Drone"), "research-global")],
  revenue=dict(unit=T("хөтөчтэй аялал", "per guided session"), attachApplies=True, capexShare=1.0, scenarios=scenarios("Thermal Drone (wildlife spotting)"))))

I.append(dict(id="night-vision", category="safety", phase=2, earns="yes", revenueCountedWith="thermal-drone",
  name=T("Шөнийн харааны төхөөрөмж", "Night-vision goggles"),
  summary=T("Ан амьтан ажиглах аялалд ашиглах шөнийн харааны төхөөрөмж. Орлогыг дронтой хамт тооцсон.", "Night-vision goggles used on the wildlife-spotting tours. Their revenue is counted with the drone."),
  components=[fixed("Night Vision Goggles", T("Шөнийн хараа", "Goggles"), "research-global")]))

shot = [fixed("Shotgun (bear defence, guide-carried)", T("Буу", "Shotgun"), "research-global"),
        fixed("Bear defence shotgun shells (stock)", T("Сумны нөөц", "Shells"), "research-global")]
shot[0]["id"], shot[0]["name"] = "shotgun", T("Буу", "Shotgun")
shot[1]["id"], shot[1]["name"] = "shells", T("Сумны нөөц", "Shells")
I.append(dict(id="bear-shotgun", category="safety", phase=1, earns="no",
  name=T("Баавгайн хамгаалалтын буу", "Bear-defence shotgun"),
  summary=T("Хөтөч аюулгүй байдлын үүднээс авч явах буу, сумны нөөцтэй. Зочдод ашиглуулахгүй.", "A shotgun and stock of bear shells, carried by guides for safety. Not for guest use."),
  components=shot))

I.append(dict(id="hunting-rifle", category="safety", phase=1, earns="no",
  name=T("Анчны буу", "Hunting rifle"),
  summary=T("Хөтөч, ан амьтны хяналтад зориулсан чанартай анчны буу.", "A quality hunting rifle for guides and wildlife management."),
  components=[fixed("Hunting Rifle", T("Буу", "Rifle"), "research-global")]))

I.append(dict(id="gun-safe", category="safety", phase=1, earns="no",
  name=T("Зэвсгийн сейф", "Gun safe"),
  summary=T("Галт зэвсгийг найдвартай хадгалах цоожтой сейф.", "A lockable long-gun safe to store the firearms securely."),
  components=[fixed("Gun Storage Safe (for your room)", T("Сейф", "Safe"), "research-global")]))

I.append(dict(id="mosquito-traps", category="safety", phase=1, earns="no",
  name=T("Шумуул барих төхөөрөмж", "Mosquito traps"),
  summary=T("Клуб байр, бөмбөгөр байшин, голын эрэг орчмын шумуулыг барьж, зуны оройг тухтай болгоно.", "Industrial traps around the clubhouse, domes and riverside to keep summer evenings comfortable."),
  components=[swap("MOSQUITO TRAP - industrial grade", opts=(
    ("Commercial propane trap - approx 1 acre coverage", "propane", T("Пропан урхи (~1 акр)", "Commercial propane trap (~1 acre)"),
     T("Тус бүр ~1 акр талбайг хамарна.", "Covers roughly 1 acre each."), "research-global", 0.20),
    ("Industrial CO2 trap - approx 2 acre coverage", "co2-industrial", T("Үйлдвэрийн CO₂ урхи (~2 акр)", "Industrial CO₂ trap (~2 acres)"),
     T("Илүү том талбай хамарна, голын эрэгт тохиромжтой.", "Larger coverage, better near the river."), "research-global", 0.20)))]))

# ---- energy
energy = [fixed("240kWh BESS Battery (in warehouse)", T("240 кВт.ц батарей", "240 kWh battery"), "research-china"),
          fixed("EV Fast Charger (60-120kW)", T("Хурдан цэнэглэгч", "Fast charger"), "research-china")]
energy[0]["id"], energy[0]["name"] = "battery", T("240 кВт.ц батарей", "240 kWh battery")
energy[1]["id"], energy[1]["name"] = "charger", T("Хурдан цэнэглэгч", "Fast charger")
I.append(dict(id="battery-ev-charging", category="energy", phase=2, earns="yes",
  name=T("Батарей ба цэнэглэгч", "Battery & EV charger"),
  summary=T("Агуулахад байрлах 240 кВт.ц батарей, цахилгаан машины хурдан цэнэглэгчтэй. Цэнэглэлтийн төлбөр олохоос гадна цахилгаан тасрахад нөөц болно.",
            "A 240 kWh battery in the warehouse paired with a fast charger for electric cars. It earns charging fees and doubles as backup power during outages."),
  components=energy,
  revenue=dict(unit=T("кВт.ц тутамд", "per kWh dispensed"), attachApplies=False, capexShare=1.0, scenarios=scenarios("240kWh BESS Battery (in warehouse)"))))

I.append(dict(id="solar-panels", category="energy", phase=2, earns="no",
  name=T("Нарны хавтан", "Solar panels"),
  summary=T("Агуулахын дээвэр дээрх ~50 кВт нарны хавтан — батарейг цэнэглэж, цахилгааны зардлыг бууруулна.", "About 50 kW of panels on the warehouse roof, charging the battery and cutting grid electricity costs."),
  components=[fixed("Solar Panel System (warehouse roof, approx 50kW)", T("Нарны систем", "Solar system"), "research-china")]))

# ---- kitchen
I.append(dict(id="wine-fridge", category="kitchen", phase=2, earns="no",
  name=T("Дарсны хөргөгч", "Wine fridge"),
  summary=T("Сонгомол хуучин дарсыг удаан хадгалах зориулалтын хөргөгч.", "A cellar-grade fridge to keep a select few vintage bottles long term."),
  components=[fixed("Wine Fridge (vintage long-term storage)", T("Хөргөгч", "Fridge"), "research-global")]))

I.append(dict(id="freezer", category="kitchen", phase=1, earns="no",
  name=T("Хөлдөөгч", "Freezers"),
  summary=T("Нэг үхрийн мах болон зуны жимсийг өвөлжүүлэх том хөлдөөгч.", "Large chest freezers to hold a cow's worth of meat plus summer berries through the winter."),
  components=[fixed("Freezer (cow + seasonal berries, in warehouse)", T("Хөлдөөгч", "Freezer"), "user-adjusted")]))

I.append(dict(id="pacojet", category="kitchen", phase=2, earns="no",
  name=T("Pacojet зайрмагны машин", "Pacojet ice-cream machine"),
  summary=T("Гал тогоонд шинэхэн зайрмаг, сорбет хийх мэргэжлийн машин.", "A professional machine for making fresh ice cream and sorbet in the kitchen."),
  components=[fixed("Pacojet 4 (ice cream machine)", T("Pacojet 4", "Pacojet 4"), "research-global")]))

I.append(dict(id="jacuzzi", category="kitchen", phase=2, earns="yes",
  name=T("Модон халуун усан сан", "Nordic wooden hot tub"),
  summary=T("6-аас дээш хүний багтаамжтай, өвөл ч халуун байх цахилгаан халаалттай модон халуун усан сан. Зочид захиалж ашиглана.",
            "A Nordic-style wooden hot tub for 6+, electrically heated to stay hot through winter. Guests book it."),
  components=[fixed("Nordic Wooden Jacuzzi (electric, 6+ capacity)", T("Халуун усан сан", "Hot tub"), "research-china")],
  revenue=dict(unit=T("нэг захиалга", "per booking"), attachApplies=True, capexShare=1.0, scenarios=scenarios("Nordic Wooden Jacuzzi (electric, 6+ capacity)"))))

# the hot tub is a guest amenity, not kitchen equipment
for it in I:
    if it["id"] == "jacuzzi": it["category"] = "amenities"

assert len(I) == 28, len(I)
for it in I: it.pop("phase", None)          # phases were dropped from the site

# ------------------------------------------------------------------ settings
settings = {
  "$schema": "./settings.schema.json",
  "lodge": {"name": "Nine Pines Lodge", "url": "https://ninepines.mn"},
  "fx": 3650,
  "guestNights": GUEST_NIGHTS,
  "defaultScenario": "cons",
  "scenarios": [
    {"id": "cons", "name": T("Болгоомжтой", "Conservative")},
    {"id": "base", "name": T("Дунд", "Base")},
    {"id": "aggr", "name": T("Өөдрөг", "Aggressive")}],
  "categories": [
    {"id": "buildings", "name": T("Барилга, дэд бүтэц", "Buildings & infrastructure")},
    {"id": "vehicles",  "name": T("Тээврийн хэрэгсэл", "Vehicles")},
    {"id": "safety",    "name": T("Аюулгүй байдал, ан амьтан", "Safety & wildlife")},
    {"id": "energy",    "name": T("Эрчим хүч", "Energy")},
    {"id": "kitchen",   "name": T("Гал тогоо", "Kitchen")},
    {"id": "amenities", "name": T("Зочдын тав тух", "Guest amenities")}],
  "sources": [
    {"id": "quote",                  "confidence": "high",   "name": T("Бодит үнийн санал", "Real quotation")},
    {"id": "given",                  "confidence": "high",   "name": T("Удирдлагаас өгсөн", "Provided by management")},
    {"id": "user-adjusted",          "confidence": "high",   "name": T("Гараар тохируулсан", "Manually set")},
    {"id": "research-mongolia",      "confidence": "medium", "name": T("Судалгаа — Монгол", "Researched — Mongolia")},
    {"id": "research-china",         "confidence": "medium", "name": T("Судалгаа — Хятадын нийлүүлэлт", "Researched — China supply")},
    {"id": "research-global",        "confidence": "medium", "name": T("Судалгаа — дэлхийн зах зээл", "Researched — global market")},
    {"id": "construction-benchmark", "confidence": "low",    "name": T("Барилгын жишиг — гүйцэтгэгчийн үнийн санал хэрэгтэй", "Construction benchmark — needs a contractor quote")}],
}

items_doc = {"$schema": "./items.schema.json", "items": I}
with open(os.path.join(ROOT, "data", "items.json"), "w", encoding="utf8") as f:
    json.dump(items_doc, f, ensure_ascii=False, indent=2)
with open(os.path.join(ROOT, "data", "settings.json"), "w", encoding="utf8") as f:
    json.dump(settings, f, ensure_ascii=False, indent=2)
print(f"wrote {len(I)} items, {sum(1 for i in I if i.get('revenue'))} with revenue, "
      f"{sum(1 for i in I for c in i['components'] if len(c['options'])>1)} swappable components")
